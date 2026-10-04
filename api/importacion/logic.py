# api/importacion/logic.py
# ─────────────────────────────────────────────────────────────────────────────
# Carga masiva de empleados desde la plantilla Excel (Fase 2).
#   leer()      → filas crudas del .xlsx, con su número de fila de Excel.
#   validar()   → por fila: crear / actualizar, errores (bloquean) y avisos
#                 (no bloquean), sin escribir nada.
#   importar()  → vuelve a validar y guarda SOLO las filas sin errores.
# Idempotente: el número de empleado identifica a la persona, así que subir el
# mismo archivo dos veces actualiza en vez de duplicar, y una celda vacía
# nunca borra un dato que ya existía.
# ─────────────────────────────────────────────────────────────────────────────
import base64
import re
import unicodedata
from datetime import date, datetime, time
from io import BytesIO

from bson.objectid import ObjectId
from openpyxl import load_workbook

from core.validadores_mx import validar_curp, validar_rfc, validar_nss, validar_clabe, normalizar, coherencia_identidad
from .columnas import COLUMNAS, CONTRATOS, REGIMENES

MAX_FILAS = 2000
MAX_BYTES = 5 * 1024 * 1024
DIAS_MES = 30.4
# Solo para un AVISO (no bloquea). Actualizar cuando se publique el vigente.
SALARIO_MINIMO_REFERENCIA = 278.80
_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


class ArchivoInvalido(Exception):
    pass


def _clave_texto(t):
    t = unicodedata.normalize("NFD", str(t or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", t.replace("*", ""))


_POR_ENCABEZADO = {_clave_texto(c[1]): c[0] for c in COLUMNAS}


def _texto(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def _fecha(v):
    """Devuelve (AAAA-MM-DD | "", error)."""
    if v in (None, ""):
        return "", None
    if isinstance(v, datetime):
        return v.date().isoformat(), None
    if isinstance(v, date):
        return v.isoformat(), None
    t = _texto(v)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(t, fmt).date().isoformat(), None
        except ValueError:
            pass
    return "", "Fecha no reconocida; usa AAAA-MM-DD o el formato de fecha de Excel."


def _hora(v):
    if v in (None, ""):
        return "", None
    if isinstance(v, (time, datetime)):
        return v.strftime("%H:%M"), None
    if isinstance(v, float) and 0 <= v < 1:   # Excel guarda horas como fracción del día
        minutos = round(v * 24 * 60)
        return f"{minutos // 60:02d}:{minutos % 60:02d}", None
    m = re.match(r"^(\d{1,2}):(\d{2})", _texto(v))
    if m and int(m.group(1)) < 24 and int(m.group(2)) < 60:
        return f"{int(m.group(1)):02d}:{m.group(2)}", None
    return "", "Hora no reconocida; usa HH:MM, ej. 09:00."


def leer(archivo_b64):
    try:
        datos = base64.b64decode(archivo_b64.split(",")[-1], validate=False)
    except Exception:
        raise ArchivoInvalido("No se pudo leer el archivo.")
    if len(datos) > MAX_BYTES:
        raise ArchivoInvalido("El archivo pesa más de 5 MB.")
    try:
        wb = load_workbook(BytesIO(datos), data_only=True, read_only=True)
    except Exception:
        raise ArchivoInvalido("El archivo no es un Excel válido (.xlsx). Descarga la plantilla y úsala como base.")
    ws = wb["Empleados"] if "Empleados" in wb.sheetnames else wb.worksheets[0]
    filas = list(ws.iter_rows(values_only=True))

    # La fila de encabezados es la que contiene "Nombre(s)" o "Nombre".
    idx_enc, mapa = None, {}
    for i, fila in enumerate(filas[:6]):
        m = {j: _POR_ENCABEZADO[_clave_texto(v)] for j, v in enumerate(fila) if _clave_texto(v) in _POR_ENCABEZADO}
        if "Nombre" in m.values() and "NumeroEmpleado" in m.values():
            idx_enc, mapa = i, m
            break
    if idx_enc is None:
        raise ArchivoInvalido("No encontré los encabezados de la plantilla (Número de empleado, Nombre…). Descarga la plantilla y úsala como base.")

    salida = []
    for i, fila in enumerate(filas[idx_enc + 1:], start=idx_enc + 2):
        valores = {mapa[j]: v for j, v in enumerate(fila) if j in mapa}
        if not any(_texto(v) for v in valores.values()):
            continue
        salida.append({"fila": i, "valores": valores})
        if len(salida) > MAX_FILAS:
            raise ArchivoInvalido(f"El archivo tiene más de {MAX_FILAS} empleados; divídelo en varios.")
    if not salida:
        raise ArchivoInvalido("La hoja «Empleados» no tiene filas con datos.")
    return salida


def _existentes(mongo):
    """Número de empleado y CURP → empleado_id de quienes ya están en el sistema."""
    por_num, por_curp = {}, {}
    for r in mongo.db.rh.find({}, {"empleado_id": 1, "NumeroEmpleado": 1, "CURP": 1}):
        eid = str(r.get("empleado_id"))
        if _texto(r.get("NumeroEmpleado")):
            por_num[_texto(r["NumeroEmpleado"]).upper()] = eid
        if _texto(r.get("CURP")):
            por_curp[normalizar(r["CURP"])] = eid
    return por_num, por_curp


def validar(mongo, filas, crear_cuentas=False):
    por_num, por_curp = _existentes(mongo)
    correos_usados = {(_texto(u.get("email")) or "").lower() for u in mongo.db.usuario.find({}, {"email": 1}) if u.get("email")}
    nums_archivo = {}
    curps_archivo = {}
    for f in filas:
        n = _texto(f["valores"].get("NumeroEmpleado")).upper()
        if n:
            nums_archivo.setdefault(n, []).append(f["fila"])
        c = normalizar(f["valores"].get("CURP"))
        if c:
            curps_archivo.setdefault(c, []).append(f["fila"])

    resultado = []
    for f in filas:
        v = f["valores"]
        errores, avisos, limpio = [], [], {}
        err = lambda campo, msg: errores.append({"campo": campo, "mensaje": msg})
        avi = lambda campo, msg: avisos.append({"campo": campo, "mensaje": msg})

        for clave in ("NumeroEmpleado", "Nombre", "ApelPaterno", "ApelMaterno", "Puesto", "Area", "JefeNumero",
                      "DiasLaborales", "Banco", "CuentaBancaria", "CorreoPersonal", "CorreoTrabajo"):
            limpio[clave] = _texto(v.get(clave))
        limpio["NumeroEmpleado"] = limpio["NumeroEmpleado"].upper()
        limpio["JefeNumero"] = limpio["JefeNumero"].upper()
        for clave, etiqueta in (("NumeroEmpleado", "Número de empleado"), ("Nombre", "Nombre"), ("ApelPaterno", "Apellido paterno")):
            if not limpio[clave]:
                err(clave, f"{etiqueta} es obligatorio.")

        num = limpio["NumeroEmpleado"]
        if num and len(nums_archivo.get(num, [])) > 1:
            err("NumeroEmpleado", f"El número {num} se repite en las filas {', '.join(map(str, nums_archivo[num]))}.")

        for clave in ("FecNacimiento", "FechaIngreso"):
            limpio[clave], e = _fecha(v.get(clave))
            if e:
                err(clave, e)
        for clave in ("HoraEntrada", "HoraSalida"):
            limpio[clave], e = _hora(v.get(clave))
            if e:
                err(clave, e)

        celular = re.sub(r"\D", "", _texto(v.get("Celular")))
        if celular and len(celular) != 10:
            avi("Celular", "El celular no tiene 10 dígitos; se guarda tal cual.")
        limpio["Celular"] = celular

        for clave, fn in (("CURP", validar_curp), ("RFC", validar_rfc), ("NSS", validar_nss), ("CLABE", validar_clabe)):
            limpio[clave] = normalizar(v.get(clave))
            if limpio[clave]:
                e = fn(limpio[clave])
                if e:
                    err(clave, e)
        curp = limpio["CURP"]
        if curp and len(curps_archivo.get(curp, [])) > 1:
            err("CURP", "La misma CURP aparece en varias filas.")

        sal = _texto(v.get("SalarioDiario")).replace("$", "").replace(",", "")
        if sal:
            try:
                s = float(sal)
                if s <= 0:
                    raise ValueError
                limpio["SalarioDiario"] = round(s, 2)
            except ValueError:
                err("SalarioDiario", "El salario diario debe ser un número mayor a 0.")
        if limpio.get("SalarioDiario") and limpio["SalarioDiario"] < SALARIO_MINIMO_REFERENCIA:
            avi("SalarioDiario", f"Está por debajo de ${SALARIO_MINIMO_REFERENCIA:.2f} diarios (salario mínimo general 2025); revísalo.")

        contrato = _texto(v.get("TipoContrato")).lower()
        if contrato:
            if contrato in CONTRATOS:
                limpio["TipoContrato"] = CONTRATOS[contrato]
            else:
                err("TipoContrato", "Contrato no reconocido; elige una opción de la lista.")
        regimen = _texto(v.get("Regimen")).lower()
        if regimen:
            if regimen in REGIMENES:
                limpio["Regimen"] = REGIMENES[regimen]
            else:
                err("Regimen", "Régimen no reconocido; usa Nómina, Asimilados u Honorarios.")

        for clave in ("CorreoPersonal", "CorreoTrabajo"):
            if limpio[clave] and not _EMAIL_RE.match(limpio[clave]):
                err(clave, "El correo no tiene un formato válido.")
        limpio["CorreoTrabajo"] = limpio["CorreoTrabajo"].lower()

        if not any(e["campo"] in ("CURP", "RFC") for e in errores):
            for campo, msg in coherencia_identidad(limpio["CURP"], limpio["RFC"], limpio["Nombre"], limpio["ApelPaterno"],
                                                   limpio["ApelMaterno"], limpio.get("FecNacimiento")).items():
                err(campo, msg)

        # ¿Crear o actualizar? Por número de empleado; si no, por CURP.
        existente = por_num.get(num) or (por_curp.get(curp) if curp else None)
        accion = "actualizar" if existente else "crear"
        if existente and curp and por_curp.get(curp) and por_curp[curp] != existente:
            err("CURP", "Esa CURP ya pertenece a otra persona del sistema.")

        jefe = limpio["JefeNumero"]
        if jefe:
            if jefe == num:
                err("JefeNumero", "Una persona no puede ser su propio jefe.")
            elif jefe not in nums_archivo and jefe not in por_num:
                err("JefeNumero", f"No encontré al jefe con número {jefe} ni en el archivo ni en el sistema.")

        if crear_cuentas and limpio["CorreoTrabajo"] and accion == "crear" and limpio["CorreoTrabajo"] in correos_usados:
            err("CorreoTrabajo", "Ese correo ya tiene una cuenta en el sistema.")
        if not limpio["CorreoTrabajo"] and crear_cuentas and accion == "crear":
            avi("CorreoTrabajo", "Sin correo de trabajo: se importa, pero sin cuenta de acceso.")

        resultado.append({
            "fila": f["fila"],
            "numero_empleado": num,
            "nombre": f"{limpio['Nombre']} {limpio['ApelPaterno']}".strip(),
            "accion": accion,
            "empleado_id": existente,
            "estado": "error" if errores else ("aviso" if avisos else "ok"),
            "errores": errores,
            "avisos": avisos,
            "_limpio": limpio,
        })
    return resultado


def _set_no_vacios(d):
    return {k: v for k, v in d.items() if v not in ("", None)}


def importar(mongo, validadas, crear_cuenta_fn=None):
    """Guarda las filas sin errores. crear_cuenta_fn(fila, empleado_id) → dict|None."""
    resumen = {"creados": 0, "actualizados": 0, "omitidos": 0, "cuentas": [], "errores_cuentas": []}
    areas = {_texto(a.get("NombreDepto")) for a in mongo.db.catalogodepto.find({}, {"NombreDepto": 1})}
    num_a_id, nombre_por_id = {}, {}
    for r in mongo.db.rh.find({}, {"empleado_id": 1, "NumeroEmpleado": 1, "NombreCompleto": 1}):
        if _texto(r.get("NumeroEmpleado")):
            num_a_id[_texto(r["NumeroEmpleado"]).upper()] = str(r.get("empleado_id"))

    pendientes_jefe = []
    for fila in validadas:
        if fila["estado"] == "error":
            resumen["omitidos"] += 1
            continue
        l = fila["_limpio"]
        area = l["Area"]
        if area and area not in areas:
            mongo.db.catalogodepto.insert_one({"NombreDepto": area, "Descripcion": "", "Poblacion": 0, "DeptoPadre": None})
            areas.add(area)
        if area and l["Puesto"]:
            # El puesto del Excel queda en el catálogo de su área para que
            # aparezca en la lista del perfil.
            mongo.db.catalogodepto.update_one({"NombreDepto": area}, {"$addToSet": {"Puestos": l["Puesto"]}})

        emp = _set_no_vacios({"Nombre": l["Nombre"], "ApelPaterno": l["ApelPaterno"], "ApelMaterno": l["ApelMaterno"],
                              "FecNacimiento": l["FecNacimiento"], "depto_id": area})
        if fila["accion"] == "crear":
            emp.setdefault("depto_id", "Sin Asignar")
            emp.update({"Fotografias": [], "Cargo": "Personal", "estado": "activo"})
            eid = mongo.db.empleados.insert_one(emp).inserted_id
            resumen["creados"] += 1
        else:
            eid = ObjectId(fila["empleado_id"])
            if emp:
                mongo.db.empleados.update_one({"_id": eid}, {"$set": emp})
            resumen["actualizados"] += 1

        horario = _set_no_vacios({"HoraEntrada": l["HoraEntrada"], "HoraSalida": l["HoraSalida"], "DiasTrabajados": l["DiasLaborales"]})
        rh = _set_no_vacios({
            "NumeroEmpleado": l["NumeroEmpleado"], "Puesto": l["Puesto"], "Departamento": area,
            "FechaIngreso": l["FechaIngreso"], "CURP": l["CURP"], "RFC": l["RFC"], "NSS": l["NSS"],
            "Banco": l["Banco"], "CLABE": l["CLABE"], "CuentaBancaria": l["CuentaBancaria"],
            "TipoRelacionLaboral": l.get("Regimen"), "tipo_contrato": l.get("TipoContrato"),
            "NombreCompleto": f"{l['Nombre']} {l['ApelPaterno']} {l['ApelMaterno']}".strip(),
        })
        if l.get("TipoContrato"):
            rh["contrato_firmado"] = l["TipoContrato"] in ("digital", "autografa")
        if l.get("SalarioDiario"):
            rh["SalarioDiario"] = l["SalarioDiario"]
            rh["Salario"] = round(l["SalarioDiario"] * DIAS_MES, 2)
        for k, v in horario.items():
            rh[f"HorarioLaboral.{k}"] = v
        mongo.db.rh.update_one({"empleado_id": eid}, {"$set": rh}, upsert=True)

        contacto = _set_no_vacios({"TelCelular": l["Celular"]})
        if l["CorreoPersonal"]:
            existente = mongo.db.datoscontacto.find_one({"EmpleadoId": eid}) or {}
            correos = existente.get("ListaCorreos") if isinstance(existente.get("ListaCorreos"), list) else []
            if l["CorreoPersonal"].lower() not in {str(c.get("email", "")).lower() for c in correos}:
                correos = correos + [{"email": l["CorreoPersonal"], "principal": not correos}]
            contacto["ListaCorreos"] = correos
        if contacto:
            mongo.db.datoscontacto.update_one({"EmpleadoId": eid}, {"$set": contacto}, upsert=True)

        num_a_id[l["NumeroEmpleado"]] = str(eid)
        nombre_por_id[str(eid)] = f"{l['Nombre']} {l['ApelPaterno']}".strip()
        if l["JefeNumero"]:
            pendientes_jefe.append((eid, l["JefeNumero"]))

        if crear_cuenta_fn and fila["accion"] == "crear" and l["CorreoTrabajo"]:
            r = crear_cuenta_fn(l, str(eid))
            (resumen["cuentas"] if r and r.get("ok") else resumen["errores_cuentas"]).append(r or {"nombre": fila["nombre"], "error": "No se pudo crear la cuenta."})

    # Segunda pasada: jefes (pueden venir después en el mismo archivo).
    for eid, jefe_num in pendientes_jefe:
        jefe_id = num_a_id.get(jefe_num)
        if not jefe_id:
            continue
        if jefe_id not in nombre_por_id:
            j = mongo.db.empleados.find_one({"_id": ObjectId(jefe_id)}, {"Nombre": 1, "ApelPaterno": 1}) or {}
            nombre_por_id[jefe_id] = f"{j.get('Nombre', '')} {j.get('ApelPaterno', '')}".strip()
        mongo.db.rh.update_one({"empleado_id": eid}, {"$set": {"JefeInmediato_id": jefe_id, "JefeInmediato": nombre_por_id[jefe_id]}})
    return resumen


def publico(validadas):
    """Quita el campo interno _limpio antes de responder."""
    return [{k: v for k, v in f.items() if k != "_limpio"} for f in validadas]
