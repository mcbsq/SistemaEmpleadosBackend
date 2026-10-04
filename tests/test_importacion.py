# tests/test_importacion.py
# Carga masiva: la plantilla se puede volver a leer, cada fila se valida con
# su motivo exacto, el jefe se resuelve aunque venga después, y subir el mismo
# archivo dos veces actualiza en vez de duplicar.
import base64
from io import BytesIO

from bson.objectid import ObjectId
from openpyxl import load_workbook

from api.importacion.logic import leer, validar, importar
from api.importacion.plantilla import generar_plantilla
from api.importacion.columnas import COLUMNAS
from tests.fakes import FakeMongo

CLAVES = [c[0] for c in COLUMNAS]


def archivo(filas):
    """Llena la plantilla real con `filas` (dicts por clave) y la devuelve en base64."""
    wb = load_workbook(BytesIO(generar_plantilla(["TECH"])))
    ws = wb["Empleados"]
    for i, f in enumerate(filas, start=3):
        for clave, valor in f.items():
            ws.cell(row=i, column=CLAVES.index(clave) + 1, value=valor)
    buf = BytesIO(); wb.save(buf)
    return base64.b64encode(buf.getvalue()).decode()


BASE = [
    {"NumeroEmpleado": "E2", "Nombre": "Luis", "ApelPaterno": "Paz", "JefeNumero": "E1", "Area": "Ventas",
     "SalarioDiario": 500, "Regimen": "Nómina", "HoraEntrada": "9:00"},
    {"NumeroEmpleado": "E1", "Nombre": "Ana", "ApelPaterno": "Ruiz", "Area": "TECH", "CorreoPersonal": "ana@mail.com"},
]


def test_plantilla_tiene_instrucciones_y_encabezados():
    wb = load_workbook(BytesIO(generar_plantilla(["TECH"])))
    assert wb.sheetnames[0] == "Instrucciones" and "Empleados" in wb.sheetnames
    assert wb["Empleados"].cell(row=2, column=1).value.startswith("Número de empleado")


def test_valida_con_motivo_exacto_por_fila():
    mongo = FakeMongo()
    filas = leer(archivo([
        {"NumeroEmpleado": "X1", "Nombre": "Sin", "ApelPaterno": "Errores"},
        {"NumeroEmpleado": "X2", "Nombre": "Mala", "ApelPaterno": "Curp", "CURP": "HEGG560427MVZRRL05", "JefeNumero": "NOEXISTE"},
        {"Nombre": "Sin número"},
    ]))
    r = validar(mongo, filas)
    assert [f["estado"] for f in r] == ["ok", "error", "error"]
    campos = {e["campo"] for e in r[1]["errores"]}
    assert {"CURP", "JefeNumero"} <= campos
    assert r[2]["fila"] == 5 and any(e["campo"] == "NumeroEmpleado" for e in r[2]["errores"])


def test_importa_resuelve_jefe_y_es_idempotente():
    mongo = FakeMongo()
    b64 = archivo(BASE)
    res = importar(mongo, validar(mongo, leer(b64)))
    assert res["creados"] == 2 and res["actualizados"] == 0
    ana = mongo.db.empleados.find_one({"Nombre": "Ana"})
    luis_rh = mongo.db.rh.find_one({"NumeroEmpleado": "E2"})
    assert luis_rh["JefeInmediato_id"] == str(ana["_id"]) and luis_rh["JefeInmediato"] == "Ana Ruiz"
    assert luis_rh["Salario"] == 15200.0 and luis_rh["TipoRelacionLaboral"] == "nomina"
    assert mongo.db.catalogodepto.find_one({"NombreDepto": "Ventas"})  # área nueva creada

    # Segunda vez: actualiza, no duplica, y una celda vacía no borra.
    res2 = importar(mongo, validar(mongo, leer(archivo([{"NumeroEmpleado": "E2", "Nombre": "Luis", "ApelPaterno": "Paz", "Puesto": "Vendedor"}]))))
    assert res2 == {**res2, "creados": 0, "actualizados": 1}
    assert len(list(mongo.db.empleados.find())) == 2
    rh = mongo.db.rh.find_one({"NumeroEmpleado": "E2"})
    assert rh["Puesto"] == "Vendedor" and rh["SalarioDiario"] == 500


def test_curp_de_otra_persona_es_error():
    mongo = FakeMongo()
    r = validar(mongo, leer(archivo([{"NumeroEmpleado": "E5", "Nombre": "Luis", "ApelPaterno": "Paz", "CURP": "HEGG560427MVZRRL04"}])))
    assert r[0]["estado"] == "error" and r[0]["errores"][0]["campo"] == "CURP"


def test_filas_con_error_no_se_guardan():
    mongo = FakeMongo()
    filas = validar(mongo, leer(archivo([{"NumeroEmpleado": "E9", "Nombre": "Mal", "ApelPaterno": "Dato", "CLABE": "123"}])))
    res = importar(mongo, filas)
    assert res["omitidos"] == 1 and res["creados"] == 0 and not list(mongo.db.empleados.find())


def test_archivo_que_no_es_excel():
    import pytest
    from api.importacion.logic import ArchivoInvalido
    with pytest.raises(ArchivoInvalido):
        leer(base64.b64encode(b"hola").decode())
